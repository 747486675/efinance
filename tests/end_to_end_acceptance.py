"""Manual real-data round trips through CLI exports and authenticated web APIs."""
import json
import os
import subprocess
import sys
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pandas as pd
import efinance as ef


def main():
    root = Path(__file__).resolve().parents[1]
    report = dict(checked_at=datetime.now(timezone(timedelta(hours=8))).isoformat(),
                  source="real upstream; closed history interval; no data fixtures", checks=[])
    old = ef.get_config()
    ef.configure(connect_timeout=5, read_timeout=10, max_attempts=1, requests_per_second=2)
    query = dict(asset="stock", codes=["1.600519"], beg="20261001", end="20261016", klt=101, fqt=1)

    def check(name, function):
        record = dict(name=name)
        try:
            details = function()
            record.update(status="PASS", **(details or {}))
            return True
        except Exception as error:
            record.update(status="FAIL", error_type=type(error).__name__, message=str(error))
            return False
        finally:
            report["checks"].append(record)
            print(record["status"], name, flush=True)

    try:
        expected = ef.stock.get_quote_history("1.600519", beg=query["beg"], end=query["end"], klt=101, fqt=1)
        if expected.empty:
            raise RuntimeError("Reference history is empty; no round trip may be counted as verified")
        report["reference"] = dict(rows=len(expected), columns=list(expected.columns), query=query)
        records = json.loads(expected.to_json(orient="records", force_ascii=False))
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            for format in ("csv", "json", "xlsx", "parquet"):
                def export(fmt=format):
                    output = folder / ("history." + fmt)
                    script = "import efinance as ef; ef.configure(connect_timeout=5,read_timeout=10,max_attempts=1,requests_per_second=2); from efinance.cli import main; import sys; sys.exit(main(sys.argv[1:]))"
                    args = [sys.executable, "-c", script, "history", "1.600519", "--beg", query["beg"], "--end", query["end"], "--format", fmt, "--output", str(output)]
                    result = subprocess.run(args, cwd=str(root), stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60, env=dict(os.environ, PYTHONIOENCODING="utf-8"))
                    if result.returncode:
                        raise RuntimeError(result.stderr.decode("utf-8", errors="replace"))
                    assert result.stdout == b"", "File export wrote data to stdout"
                    if fmt == "csv":
                        assert output.read_bytes().startswith(b"\xef\xbb\xbf")
                        assert output.read_bytes().decode("utf-8-sig") == expected.to_csv(index=False)
                    elif fmt == "json":
                        assert json.loads(output.read_text(encoding="utf-8")) == records
                    elif fmt == "xlsx":
                        from openpyxl import load_workbook
                        book = load_workbook(output)
                        try:
                            column = list(expected.columns).index("股票代码") + 1
                            assert book.active.cell(2, column).data_type == "s"
                        finally:
                            book.close()
                        pd.testing.assert_frame_equal(pd.read_excel(output, dtype={"股票代码": str}), expected, check_dtype=False)
                    else:
                        pd.testing.assert_frame_equal(pd.read_parquet(output), expected)
                    before = output.read_bytes()
                    rerun = subprocess.run(args, cwd=str(root), stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=15)
                    assert rerun.returncode == 2, "Overwrite protection did not return parameter error"
                    assert output.read_bytes() == before, "Overwrite protection changed existing bytes"
                    return dict(rows=len(expected), bytes=len(before), overwrite_protection="PASS")
                check("cli." + format, export)

        from fastapi.testclient import TestClient
        from efinance.web import create_app
        with patch.dict(os.environ, {"EFINANCE_WEB_USERNAME": "acceptance-local", "EFINANCE_WEB_PASSWORD": "temporary-test-only"}), TestClient(create_app()) as client:
            def health():
                assert client.get("/api/health").status_code == 401
                response = client.get("/api/health", auth=("acceptance-local", "temporary-test-only"))
                assert response.status_code == 200 and response.json()["status"] == "ok"
            check("web.auth_and_health", health)
            client.auth = ("acceptance-local", "temporary-test-only")

            def history():
                response = client.post("/api/history", json=query)
                assert response.status_code == 200, response.text
                value = response.json()["results"][0]
                assert value["columns"] == list(expected.columns)
                assert value["records"] == records
                return dict(rows=len(value["records"]))
            check("web.history_equals_python", history)

            def download():
                response = client.post("/api/export", json=dict(query, kind="history"))
                assert response.status_code == 200, response.text
                assert "attachment" in response.headers["content-disposition"]
                assert response.content.startswith(b"\xef\xbb\xbf")
                frame = expected.copy()
                frame.insert(0, "查询标的", "1.600519")
                assert response.content.decode("utf-8-sig") == frame.to_csv(index=False)
                return dict(rows=len(frame), bytes=len(response.content))
            check("web.csv_equals_python", download)
    except Exception as error:
        report["checks"].append(dict(name="acceptance_setup", status="FAIL", error_type=type(error).__name__, message=str(error)))
    finally:
        ef.configure(**old)
        report["summary"] = {status: sum(r["status"] == status for r in report["checks"])
                             for status in ("PASS", "FAIL")}
        Path("docs/end-to-end-acceptance-results.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report["summary"]))
    return int(report["summary"]["FAIL"] > 0)


if __name__ == "__main__":
    sys.exit(main())
